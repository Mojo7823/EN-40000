"""Risk Management Elements section builder."""
from typing import Iterable, Optional
import re

from docx import Document
from docx.shared import Pt

from .html_converter import append_html_to_document

EVIDENCE_STATUS_LABELS = {
    "complete": "Complete",
    "in_progress": "In Progress",
    "not_started": "Not Started",
}


def append_risk_management_section(document: Document, payload: Optional[object]) -> None:
    """Append Section 5 - Risk Management Elements to the document."""
    if not payload:
        return

    general_html = _extract_value(payload, "general_approach_html")
    product_context_payload = getattr(payload, "product_context", None)
    has_product_context = _product_context_has_content(product_context_payload)

    if not general_html and not has_product_context:
        return

    document.add_page_break()

    heading = document.add_paragraph()
    heading_run = heading.add_run("5. RISK MANAGEMENT ELEMENTS")
    heading_run.font.size = Pt(20)
    heading_run.font.bold = True
    heading.space_after = Pt(8)

    reference = document.add_paragraph("[Reference: Clause 6.1 - General]")
    reference.runs[0].font.bold = True
    reference.space_after = Pt(10)

    if general_html:
        subheading = document.add_paragraph()
        subheading_run = subheading.add_run("5.1 General Approach to Risk Management")
        subheading_run.font.size = Pt(18)
        subheading_run.font.bold = True
        subheading.space_before = Pt(6)
        subheading.space_after = Pt(6)

        clause_reference = document.add_paragraph("[Reference: Clause 5 - Risk Management Elements]")
        clause_reference.runs[0].font.bold = True
        clause_reference.space_after = Pt(10)

        append_html_to_document(document, general_html)

    if has_product_context:
        _append_product_context_section(document, product_context_payload)


def _append_product_context_section(document: Document, payload: Optional[object]) -> None:
    heading = document.add_paragraph()
    heading_run = heading.add_run("5.2 Product Context")
    heading_run.font.size = Pt(18)
    heading_run.font.bold = True
    heading.space_before = Pt(12)
    heading.space_after = Pt(6)

    reference = document.add_paragraph("[Reference: Clause 6.2 - Product Context]")
    reference.runs[0].font.bold = True
    reference.space_after = Pt(10)

    subsection = document.add_paragraph()
    subsection_run = subsection.add_run("5.2.1 Intended Purpose and Reasonably Foreseeable Use")
    subsection_run.font.size = Pt(16)
    subsection_run.font.bold = True
    subsection.space_before = Pt(6)
    subsection.space_after = Pt(4)

    clause_reference = document.add_paragraph(
        "[Reference: Clause 6.2.1.2 - Product intended purpose and reasonable foreseeable use]"
    )
    clause_reference.runs[0].font.bold = True
    clause_reference.space_after = Pt(8)

    intro = document.add_paragraph(
        "The product context provides the foundation for all risk management activities. It describes the product's intended purpose, functions, operational environment, architecture, and users."
    )
    intro.paragraph_format.space_after = Pt(12)

    _append_product_context_block(
        document,
        "Intended Purpose",
        _extract_value(payload, "intended_purpose_html"),
    )
    _append_product_context_block(
        document,
        "Specific Intended Uses",
        _extract_value(payload, "specific_intended_uses_html"),
    )
    _append_product_context_block(
        document,
        "Reasonably Foreseeable Use and Misuse Considerations",
        _extract_value(payload, "foreseeable_use_html"),
    )

    _append_evidence_tracker(document, getattr(payload, "evidence_entries", None))


def _append_product_context_block(document: Document, title: str, html_content: Optional[str]) -> None:
    if not html_content:
        return
    subheading = document.add_paragraph()
    subheading_run = subheading.add_run(title)
    subheading_run.font.size = Pt(13)
    subheading_run.font.bold = True
    subheading.space_before = Pt(6)
    subheading.space_after = Pt(4)
    append_html_to_document(document, html_content)


def _append_evidence_tracker(document: Document, payload: Optional[Iterable[object]]) -> None:
    entries = _normalize_evidence_entries(payload)
    if not entries:
        return

    heading = document.add_paragraph()
    heading_run = heading.add_run("Evidence Tracker")
    heading_run.font.size = Pt(13)
    heading_run.font.bold = True
    heading.space_before = Pt(10)
    heading.space_after = Pt(4)

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Evidence Reference", "Title / Artifact", "Status", "Notes"]
    header_cells = table.rows[0].cells
    for idx, label in enumerate(headers):
        paragraph = header_cells[idx].paragraphs[0]
        run = paragraph.add_run(label)
        run.font.bold = True

    for entry in entries:
        row = table.add_row().cells
        row[0].text = entry["reference"] or ""
        row[1].text = entry["title"] or ""
        row[2].text = entry["status_label"]
        row[3].text = entry["notes"] or ""


def _normalize_evidence_entries(payload: Optional[Iterable[object]]):
    if not payload:
        return []
    normalized = []
    for entry in payload:
        reference = _extract_value(entry, "reference_id")
        title = _extract_value(entry, "title")
        status = _extract_value(entry, "status") or "not_started"
        notes = _strip_html(_extract_value(entry, "notes_html"))
        if not any([reference, title, notes]):
            continue
        normalized.append(
            {
                "reference": reference,
                "title": title,
                "status_label": EVIDENCE_STATUS_LABELS.get(status, status.replace("_", " ").title()),
                "notes": notes,
            }
        )
    return normalized


def _product_context_has_content(payload: Optional[object]) -> bool:
    if not payload:
        return False
    return any(
        [
            bool(_extract_value(payload, "intended_purpose_html")),
            bool(_extract_value(payload, "specific_intended_uses_html")),
            bool(_extract_value(payload, "foreseeable_use_html")),
            bool(_normalize_evidence_entries(getattr(payload, "evidence_entries", None))),
        ]
    )


def _extract_value(payload: object, field: str) -> Optional[str]:
    if isinstance(payload, dict):
        value = payload.get(field)
    else:
        value = getattr(payload, field, None)
    if value and isinstance(value, str) and value.strip():
        return value
    return None


def _strip_html(value: Optional[str]) -> str:
    if not value:
        return ""
    return re.sub(r"<[^>]+>", " ", value).strip()
